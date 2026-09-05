import asyncio
import csv
import io
import json
import socket
from contextlib import closing

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import database as db, main, security
from app.crawler import CrawlEngine, canonicalize_url
from app.extractor import extract_records, discover_links, DATE_RE, PHONE_RE
from app.normalizer import entity_key, record_hash
from app.runtime import single_instance
from fixture_site import DirectorySite


@pytest.mark.parametrize("url,expected", [
    ("https://EXAMPLE.com:443/a?b=2&a=1&utm_source=x#top", "https://example.com/a?a=1&b=2"),
    ("http://Example.com:80", "http://example.com/"),
    ("https://example.com/?ref=2&id=3&id=1", "https://example.com/?id=3&id=1&ref=2"),
    ("https://example.com/?empty=&q=a%26b", "https://example.com/?empty=&q=a%26b"),
    ("https://example.com:bad/", ""), ("https://user:pass@example.com/", ""),
    ("javascript:alert(1)", ""), ("http:///missing", ""),
    ("https://example.com/\n", ""), ("http://[::1]/", "http://[::1]/"),
])
def test_urls(url, expected):
    assert canonicalize_url(url) == expected


@pytest.mark.parametrize("url", ["https://other.test/", "https://fixture.test.evil.test/", "http://fixture.test:8080/", "https://fixture.test/x.pdf"])
async def test_domain_boundary(source, url):
    assert not CrawlEngine(source, 0)._allowed_url(url)


def test_identity_and_hash():
    a = {"name": "Jane Doe", "source_url": "https://x/a", "phone": "6085551111"}
    assert entity_key(a) != entity_key({**a, "source_url": "https://x/b"})
    assert entity_key(a) == entity_key({**a, "phone": "6085552222"})
    assert record_hash(a) == record_hash({**a, "extra": {"raw_text": "changed navigation"}})
    assert record_hash(a) != record_hash({**a, "phone": "6085552222"})
    assert canonicalize_url("https://x/a/../b/%7e?q=1") == "https://x/b/~?q=1"


async def test_duplicate_changes_and_history(source):
    a = {"external_id": "P1", "name": "Jane", "phone": "6085551111", "source_url": "https://fixture.test/jane"}
    results = await asyncio.gather(*(db.upsert_record(source["id"], a) for _ in range(12)))
    assert results.count("new") == 1
    assert results.count("unchanged") == 11
    assert await db.upsert_record(source["id"], {**a, "phone": "6085552222"}) == "updated"
    rows = await db.search_records()
    assert rows["total"] == 1
    history = await db.record_history(rows["items"][0]["id"])
    assert len(history) == 2
    assert json.loads(history[0]["old_json"])["phone"] == "608-555-1111"
    assert json.loads(history[0]["new_json"])["phone"] == "608-555-2222"
    assert (await db.search_records(q="P1"))["total"] == 1
    assert (await db.search_records(q="%"))["total"] == 0


async def test_two_incremental_scans(source):
    site = DirectorySite()
    first = await db.create_job(source["id"], False)
    await CrawlEngine(source, first, transport=httpx.MockTransport(site)).run()
    assert (await db.get_job(first))["status"] == "completed"
    assert (await db.search_records())["total"] == 2
    before = (await db.search_records(q="John"))["items"][0]["last_seen"]
    site.version = 2
    second = await db.create_job(source["id"], False)
    await CrawlEngine(source, second, transport=httpx.MockTransport(site)).run()
    job = await db.get_job(second)
    assert job["status"] == "completed"
    assert job["records_new"] == 1 and job["records_updated"] == 1
    assert (await db.search_records())["total"] == 3
    assert (await db.search_records(q="John"))["items"][0]["last_seen"] > before
    assert any(path == "/" and "if-none-match" in headers for path, headers in site.requests)


async def test_inactive_only_when_complete(source):
    record = {"name": "Gone", "external_id": "gone", "source_url": "https://fixture.test/gone"}
    await db.upsert_record(source["id"], record)
    site = DirectorySite()
    limited = {**source, "max_pages": 1}
    job = await db.create_job(source["id"], False)
    await CrawlEngine(limited, job, transport=httpx.MockTransport(site)).run()
    assert (await db.get_job(job))["status"] == "partial"
    assert (await db.search_records(q="Gone"))["items"][0]["active"] == 1
    job = await db.create_job(source["id"], False)
    await CrawlEngine(source, job, transport=httpx.MockTransport(site)).run()
    assert (await db.search_records(q="Gone"))["items"][0]["active"] == 0
    assert (await db.search_records(q="Gone"))["items"][0]["inactive_since"]


async def test_migration_preserves_history(source):
    record = {"name": "Legacy", "source_url": "https://fixture.test/person", "phone": "6085551111"}
    await db.upsert_record(source["id"], record)
    await db.upsert_page(source["id"], source["start_url"], content_hash="old", etag="old")
    with closing(db.connect()) as conn:
        conn.execute("UPDATE records SET entity_key='legacy-key'")
        conn.execute("PRAGMA user_version=1")
        conn.commit()
    await db.init_db()
    rows = await db.search_records()
    assert rows["total"] == 1
    assert rows["items"][0]["entity_key"] == entity_key(record)
    assert len(await db.record_history(rows["items"][0]["id"])) == 1
    assert (await db.get_page(source["id"], source["start_url"]))["content_hash"] is None


async def test_conflicting_duplicates_are_reported(source):
    site = DirectorySite()
    def response(request):
        if request.url.path == "/page/2":
            return httpx.Response(200, text=site.profile("P1", "Jane Doe", "6085559999"))
        return site(request)
    job = await db.create_job(source["id"], False)
    await CrawlEngine(source, job, transport=httpx.MockTransport(response)).run()
    result = await db.get_job(job)
    assert result["status"] == "partial"
    assert result["errors"] == 1


async def test_cached_duplicate_conflict(source):
    site = DirectorySite()
    def response(request):
        if request.url.path == "/page/2":
            if "if-none-match" in request.headers:
                return httpx.Response(304)
            return httpx.Response(200, text=site.profile("P1", "Jane Doe", "608-555-1111"), headers={"etag": "duplicate"})
        return site(request)
    job = await db.create_job(source["id"], False)
    await CrawlEngine(source, job, transport=httpx.MockTransport(response)).run()
    assert (await db.get_job(job))["status"] == "completed"
    site.version = 2
    job = await db.create_job(source["id"], False)
    await CrawlEngine(source, job, transport=httpx.MockTransport(response)).run()
    assert (await db.get_job(job))["status"] == "partial"


async def test_robots_excludes_page(source):
    seen = []
    def response(request):
        seen.append(request.url.path)
        if request.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nDisallow: /private")
        if request.url.path == "/sitemap.xml":
            return httpx.Response(404)
        return httpx.Response(200, text='<a href="/private">private</a>')
    job = await db.create_job(source["id"], False)
    await CrawlEngine(source, job, transport=httpx.MockTransport(response)).run()
    assert "/private" not in seen
    assert (await db.get_job(job))["status"] == "partial"


async def test_sitemap_direct_children_only_and_encoding(source):
    def response(request):
        body = '<?xml version="1.0" encoding="ISO-8859-1"?><urlset xmlns:image="urn:image"><url><loc>https://fixture.test/caf&#233;</loc><image:image><image:loc>https://fixture.test/ignored</image:loc></image:image></url></urlset>'
        return httpx.Response(200, content=body.encode("iso-8859-1"), headers={"content-type": "application/xml"})
    job = await db.create_job(source["id"], False)
    engine = CrawlEngine(source, job)
    async with httpx.AsyncClient(transport=httpx.MockTransport(response)) as client:
        urls = await engine._discover_sitemap_urls(client)
    assert urls == ["https://fixture.test/café"]


async def test_scan_launch_is_serialized(source, monkeypatch):
    entered = asyncio.Event()
    release = asyncio.Event()
    class SlowEngine:
        def __init__(self, *args, **kwargs):
            pass
        async def run(self):
            entered.set()
            await release.wait()
    monkeypatch.setattr(main, "CrawlEngine", SlowEngine)
    jobs = await asyncio.gather(*(main.launch_scan(source["id"]) for _ in range(8)))
    await entered.wait()
    assert len({j["id"] for j in jobs}) == 1
    tasks = list(main.TASKS.values())
    release.set()
    await asyncio.gather(*tasks)


async def test_per_host_throttle_shared(source):
    a = CrawlEngine({**source, "delay_ms": 20}, 0)
    b = CrawlEngine({**source, "delay_ms": 20}, 0)
    import time
    await a._throttle()
    start = time.monotonic()
    await b._throttle()
    assert time.monotonic() - start >= 0.015


@pytest.mark.parametrize("max_pages,max_depth", [(1, 12), (2, 12), (100, 0), (100, 1)])
async def test_limits_and_depth(source, max_pages, max_depth):
    fetched = []
    def site(request):
        path = request.url.path
        if path in {"/robots.txt", "/sitemap.xml"}:
            return httpx.Response(404)
        fetched.append(path)
        return httpx.Response(200, text=f'<a href="/{len(path)}x">next</a>')
    configured = {**source, "max_pages": max_pages, "max_depth": max_depth, "concurrency": 32}
    job = await db.create_job(source["id"], False)
    await CrawlEngine(configured, job, transport=httpx.MockTransport(site)).run()
    assert len(fetched) <= max_pages
    assert len(fetched) <= max_depth + 1
    assert (await db.get_job(job))["pages_processed"] == len(fetched)


async def test_shortest_depth_with_slow_parent(source):
    fetched = []
    async def site(request):
        path = request.url.path
        if path in {"/robots.txt", "/sitemap.xml"}:
            return httpx.Response(404)
        fetched.append(path)
        if path == "/slow":
            await asyncio.sleep(0.02)
        pages = {"/": '<a href="/fast">F</a><a href="/slow">S</a>',
                 "/fast": '<a href="/middle">M</a>', "/middle": '<a href="/target">T</a>',
                 "/slow": '<a href="/target">T</a>', "/target": '<a href="/child">C</a>', "/child": ""}
        return httpx.Response(200, text=pages[path])
    job = await db.create_job(source["id"], False)
    await CrawlEngine({**source, "max_depth": 3}, job, transport=httpx.MockTransport(site)).run()
    assert "/child" in fetched


async def test_redirect_blocked_before_request(source):
    seen = []
    def site(request):
        seen.append(str(request.url))
        return httpx.Response(302, headers={"location": "http://127.0.0.1/"})
    async with httpx.AsyncClient(transport=httpx.MockTransport(site)) as client:
        with pytest.raises(ValueError, match="boundary"):
            await CrawlEngine(source, 0)._http_fetch_with_retry(client, source["start_url"], {})
    assert len(seen) == 1


@pytest.mark.parametrize("address", ["127.0.0.1", "10.0.0.1", "172.16.0.1", "192.168.1.1", "169.254.169.254", "::1", "fc00::1", "224.0.0.1", "::ffff:127.0.0.1", "100.64.0.1"])
def test_private_addresses(address):
    assert not security.public_address(address)


async def test_dns_mixed_answers_blocked(monkeypatch):
    async def resolve(*args, **kwargs):
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 80)),
                (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("127.0.0.1", 80))]
    monkeypatch.setattr(asyncio.get_running_loop(), "getaddrinfo", resolve)
    with pytest.raises(ValueError, match="non-public"):
        await security.validate_public_url("http://example.test/")


async def test_transport_pins_dns_and_tls(monkeypatch):
    calls = []
    async def resolve(url):
        calls.append(url)
        return ["93.184.216.34"]
    async def send(self, request):
        assert request.url.host == "93.184.216.34"
        assert request.headers["host"] == "example.test"
        assert request.extensions["sni_hostname"] == "example.test"
        return httpx.Response(200, text="ok")
    monkeypatch.setattr(security, "validate_public_url", resolve)
    monkeypatch.setattr(httpx.AsyncHTTPTransport, "handle_async_request", send)
    async with httpx.AsyncClient(transport=security.PublicTransport()) as client:
        assert (await client.get("https://example.test/")).text == "ok"
    assert calls == ["https://example.test/"]


async def test_robots_fail_closed(source):
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(403))) as client:
        with pytest.raises(ValueError, match="robots"):
            await CrawlEngine(source, 0)._load_robots(client)


async def test_challenge_never_rendered(source, monkeypatch):
    engine = CrawlEngine({**source, "render_mode": "auto"}, 0)
    async def no_render(*args):
        pytest.fail("Challenge was sent to renderer")
    monkeypatch.setattr(engine.renderer, "fetch", no_render)
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(200, text="<script></script>Verify you are human"))) as client:
        result = await engine._fetch(client, source["start_url"], None)
        assert result.status == 200


async def test_oversized_response(source, monkeypatch):
    import app.crawler as crawler
    monkeypatch.setattr(crawler, "MAX_BODY_BYTES", 10)
    async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(200, content=b"x" * 11))) as client:
        with pytest.raises(ValueError, match="body limit"):
            await CrawlEngine(source, 0)._http_fetch_with_retry(client, source["start_url"], {})


def test_extraction_safety_and_links():
    assert not extract_records("<h1>Contact us</h1><footer>608-555-1111</footer>", "https://x/")
    assert DATE_RE.search("Filed 09/05/2026").group() == "09/05/2026"
    assert DATE_RE.search("Filed 2026-09-05").group() == "2026-09-05"
    assert not PHONE_RE.search("123456789012345")
    assert discover_links('<base href="/dir/"><a href="person">P</a><link rel="next" href="?page=2">', "https://x/") == ["https://x/dir/person", "https://x/dir/?page=2"]
    html = '<table><tr><th>Name</th><th>Phone</th></tr><tr><td><a href="javascript:alert(1)">Jane</a></td><td>6085551111</td></tr></table>'
    assert extract_records(html, "https://x/")[0]["source_url"] == "https://x/"


def test_single_instance_lock(database):
    with single_instance():
        with pytest.raises(RuntimeError, match="another process"):
            with single_instance():
                pass


async def test_restart_recovery(source):
    job = await db.create_job(source["id"], False)
    await db.mark_interrupted_jobs()
    assert (await db.get_job(job))["status"] == "interrupted"


def test_api_smoke_and_exports(database, monkeypatch):
    site = DirectorySite()
    class FixtureEngine(CrawlEngine):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs, transport=httpx.MockTransport(site))
    async def public(url):
        return ["93.184.216.34"]
    monkeypatch.setattr(main, "CrawlEngine", FixtureEngine)
    monkeypatch.setattr(main, "validate_public_url", public)
    with TestClient(main.app) as client:
        assert client.get("/").status_code == 200
        assert client.get("/api/health").json()["ok"]
        assert "script-src 'self'" in client.get("/").headers["content-security-policy"]
        assert client.post("/api/sources", json={"name": "Bad"}, headers={"origin": "https://evil.test"}).status_code == 403
        assert client.get("/", headers={"host": "evil.test"}).status_code == 400
        result = client.post("/api/sources", json={"name": "Fixture", "start_url": "https://fixture.test/", "delay_ms": 0, "render_mode": "http"})
        assert result.status_code == 201, result.text
        sid = result.json()["id"]
        assert len(client.get("/api/sources").json()) == 1
        assert client.post("/api/sources", json={"name": "Duplicate", "start_url": "https://fixture.test/"}).status_code == 409
        for version in (1, 2):
            site.version = version
            job = client.post(f"/api/sources/{sid}/scan", json={}).json()
            import time
            for _ in range(200):
                status = client.get(f"/api/jobs/{job['id']}").json()
                if status["status"] not in {"queued", "running"}:
                    break
                time.sleep(0.01)
            assert status["status"] == "completed", status
        assert client.get("/api/records?q=Jane").json()["total"] == 1
        assert client.get("/api/records").json()["total"] == 3
        with closing(db.connect()) as conn:
            conn.execute("UPDATE records SET name='=1+1' WHERE external_id='P1'")
            conn.commit()
        csv_response = client.get("/api/export?format=csv")
        assert csv_response.status_code == 200
        rows = list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
        assert any(row["name"] == "'=1+1" for row in rows)
        xlsx = client.get("/api/export?format=xlsx")
        book = load_workbook(io.BytesIO(xlsx.content))
        assert book.active.max_row == 4
        assert all(cell.data_type != "f" for row in book.active for cell in row)
        assert any(row["name"] == "=1+1" for row in client.get("/api/export?format=json").json())
        assert client.get("/api/export?format=bad").status_code == 400
