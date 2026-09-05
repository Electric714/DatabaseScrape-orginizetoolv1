import csv
import io
import json
from contextlib import closing

import httpx
import pytest
from openpyxl import load_workbook

from app import database as db, main
from app.extractor import extract_records
from app.normalizer import canonical_record, normalize_osha_status, record_hash


@pytest.mark.parametrize("raw,expected", [
    (None, "unknown"), ("", "unknown"), ("Open", "open"), ("closed", "closed"),
    ("No open violations reported", "none_reported"), ("0", "unknown"),
    ("none", "unknown"), ("no record found", "unknown"),
    ("An open violation was reported in 2016", "unknown"),
    ("Closed with an open appeal", "unknown"),
])
def test_osha_status_requires_exact_source_status(raw, expected):
    assert normalize_osha_status(raw) == expected


def test_labeled_business_owner_location_and_osha_are_distinct():
    html = """<table><tr><th>Person name</th><th>Business name</th><th>Owner name</th>
      <th>Business address</th><th>Location</th><th>OSHA status</th><th>OSHA violations</th></tr>
      <tr><td>Amy Contact</td><td>Fixture Builders LLC</td><td>Pat Owner</td><td>12 Oak Rd</td>
      <td>Madison, WI</td><td>Open</td><td>Inspection 123: citation under review</td></tr></table>"""
    record, = extract_records(html, "https://fixture.test/business")
    assert (record["name"], record["company"], record["owner"]) == ("Amy Contact", "Fixture Builders LLC", "Pat Owner")
    assert (record["address"], record["location"]) == ("12 Oak Rd", "Madison, WI")
    assert record["osha_status"] == "open"
    assert record["osha_details"] == "Inspection 123: citation under review"
    assert record["extra"]["osha_status_raw"] == "Open"
    assert record["extra"]["collected_from_url"] == "https://fixture.test/business"
    different_provenance = {**record, "extra": {**record["extra"], "collected_from_url": "https://fixture.test/another"}}
    assert record_hash(record) == record_hash(different_provenance)


def test_labeled_cards_and_definition_lists_do_not_invent_owners_or_osha():
    html = """<article><h2>Research result</h2><p>Business name: One LLC</p>
      <p><strong>Owner:</strong> Owner One</p><p>Location: Milwaukee, WI</p>
      <p>OSHA violations: A historical open citation is mentioned here.</p></article>
      <dl><dt>Business name</dt><dd>Two LLC</dd><dt>Name</dt><dd>Contact Two</dd>
      <dt>Address</dt><dd>20 Main St</dd></dl>"""
    rows = {r["company"]: r for r in extract_records(html, "https://fixture.test/list")}
    assert rows["One LLC"]["owner"] == "Owner One"
    assert rows["One LLC"]["name"] == ""
    assert rows["One LLC"]["location"] == "Milwaukee, WI"
    assert rows["One LLC"]["osha_status"] == "unknown"
    assert rows["One LLC"]["osha_details"].startswith("A historical open citation")
    assert rows["Two LLC"]["name"] == "Contact Two"
    assert rows["Two LLC"]["owner"] == ""
    assert rows["Two LLC"]["osha_status"] == "unknown"
    assert rows["Two LLC"]["osha_details"] == ""


async def test_research_migration_preserves_records_and_original_history(source):
    record = {"external_id": "old-1", "name": "Legacy Contact", "company": "Legacy LLC",
              "source_url": "https://fixture.test/legacy", "address": "1 Main St"}
    await db.upsert_record(source["id"], record)
    before = (await db.search_records())["items"][0]
    history = await db.record_history(before["id"])
    await db.upsert_page(source["id"], source["start_url"], content_hash="cached", etag="old")
    # Recreate the actual prior schema, including absent research columns.
    with closing(db.connect()) as conn:
        conn.execute("DROP INDEX idx_records_osha")
        for field in ("owner", "location", "osha_details", "osha_status"):
            conn.execute(f"ALTER TABLE records DROP COLUMN {field}")
        conn.execute("UPDATE records SET content_hash='previous-format-hash'")
        conn.execute("PRAGMA user_version=2")
        conn.commit()
    await db.init_db()
    after = await db.get_record(before["id"])
    assert after["name"] == before["name"] and after["owner"] == ""
    assert after["osha_status"] == "unknown" and after["osha_details"] == ""
    assert after["first_seen"] == before["first_seen"]
    assert after["last_changed"] == before["last_changed"]
    assert await db.record_history(before["id"]) == history
    assert after["content_hash"] == record_hash(record)
    assert await db.upsert_record(source["id"], record) == "unchanged"
    assert (await db.get_page(source["id"], source["start_url"]))["etag"] is None
    await db.init_db()  # Startup migration is idempotent.
    assert await db.record_history(before["id"]) == history


async def test_research_updates_preserve_source_wording_in_history(source):
    record = {"external_id": "business-1", "company": "Builders LLC", "owner": "Original Owner",
              "source_url": "https://fixture.test/business", "osha_status": "Open",
              "osha_details": "Inspection 123; source says open"}
    assert await db.upsert_record(source["id"], record) == "new"
    changed = {**record, "owner": "New Owner", "location": "Milwaukee, WI", "osha_status": "Closed",
               "osha_details": "Inspection 123; closed on source"}
    assert await db.upsert_record(source["id"], changed) == "updated"
    row = (await db.search_records())["items"][0]
    assert row["owner"] == "New Owner" and row["osha_status"] == "closed"
    history = await db.record_history(row["id"])
    assert len(history) == 2
    assert json.loads(history[0]["old_json"])["owner"] == "Original Owner"
    assert json.loads(history[0]["new_json"])["osha_details"] == changed["osha_details"]
    assert canonical_record({"company": "No details LLC"})["osha_status"] == "unknown"


async def test_research_api_filters_details_and_matching_safe_exports(source):
    records = [
        {"external_id": "business-1", "company": "Builders LLC", "name": "Jordan Contact", "owner": "=OwnerName()",
         "address": "10 Oak Rd", "location": "Madison, WI", "osha_status": "Open",
         "osha_details": "+Citation 123 source wording", "source_url": "https://fixture.test/one"},
        {"external_id": "business-2", "company": "Jordan Company", "owner": "Different Owner",
         "address": "20 Oak Rd", "location": "Milwaukee, WI", "source_url": "https://fixture.test/two"},
        {"external_id": "business-3", "company": "Closed LLC", "owner": "Jordan Owner",
         "osha_status": "Closed", "source_url": "https://fixture.test/three"},
    ]
    for record in records:
        await db.upsert_record(source["id"], record)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://testserver") as client:
        assert (await client.get("/api/records", params={"q": "Jordan"})).json()["total"] == 3
        for field, expected in (("company", "business-2"), ("name", "business-1"), ("owner", "business-3")):
            result = (await client.get("/api/records", params={"q": "Jordan", "field": field})).json()
            assert result["total"] == 1 and result["items"][0]["external_id"] == expected
        assert (await client.get("/api/records", params={"q": "Madison", "field": "location"})).json()["total"] == 1
        assert (await client.get("/api/records", params={"q": "Citation 123"})).json()["total"] == 1
        assert (await client.get("/api/records", params={"q": "business-2"})).json()["total"] == 1
        assert (await client.get("/api/records", params={"q": "%"})).json()["total"] == 0
        assert (await client.get("/api/records", params={"osha_status": "unknown"})).json()["total"] == 1
        query = {"osha_status": "open", "field": "owner", "q": "OwnerName", "source_id": source["id"]}
        rows = (await client.get("/api/records", params=query)).json()
        assert rows["total"] == 1
        rid = rows["items"][0]["id"]
        detail = (await client.get(f"/api/records/{rid}")).json()
        assert detail["source_name"] == source["name"]
        assert detail["source_url"] == "https://fixture.test/one"
        assert all(detail[key] for key in ("first_seen", "last_seen", "last_changed", "active"))
        assert len((await client.get(f"/api/records/{rid}/history")).json()) == 1
        assert (await client.get("/api/records/999999")).status_code == 404
        for endpoint in ("/api/records", "/api/export"):
            assert (await client.get(endpoint, params={"field": "owner); DROP TABLE records;"})).status_code == 422
            assert (await client.get(endpoint, params={"osha_status": "safe"})).status_code == 422
        exported = await client.get("/api/export", params={**query, "format": "json"})
        assert len(exported.json()) == 1
        assert exported.json()[0]["owner"] == "=OwnerName()"
        assert exported.json()[0]["osha_details"] == "+Citation 123 source wording"
        response = await client.get("/api/export", params={**query, "format": "csv"})
        csv_rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
        assert len(csv_rows) == 1 and csv_rows[0]["owner"] == "'=OwnerName()"
        assert csv_rows[0]["osha_details"] == "'+Citation 123 source wording"
        response = await client.get("/api/export", params={**query, "format": "xlsx"})
        book = load_workbook(io.BytesIO(response.content))
        assert book.active.max_row == 2
        assert all(cell.data_type != "f" for row in book.active for cell in row)
